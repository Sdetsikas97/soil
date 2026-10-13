#!/usr/bin/env python3
"""Rebuild the atlas data files from the review spreadsheet.

Usage (from the repository root):
    pip install openpyxl
    python scripts/build_data.py                       # uses data/source/soil_spectral_libraries_structured.xlsx
    python scripts/build_data.py path/to/workbook.xlsx # or any workbook with the same sheets

Reads the sheets "Papers", "Libraries" and "Library countries" and writes:
    data/libraries.json      one record per library (+ point location for mapping)
    data/papers.json         every screened paper with its decision and library links
    data/countries.geojson   country polygons with SSL coverage attributes (for GIS download)
    data/meta.json           counts and build date

The screening decision is recomputed here from the coded columns (C1-C3, duplicate),
so the site stays correct even if the workbook was saved without recalculating formulas.
"""
import json, math, re, sys, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/soil_spectral_libraries_structured.xlsx"
WORLD = ROOT / "data/source/world_50m.geojson"
OUT = ROOT / "data"

# Spreadsheet country names -> Natural Earth names used in the boundary file
ALIAS = {"United States": "United States of America", "DR Congo": "Dem. Rep. Congo"}
SCALE_RANK = {"National": 0, "Sub-national": 1, "Multi-country": 2}


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s in ("", "-", "NA", "None", "Not recorded") else s


def header_map(ws):
    return {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value is not None}


def spectral_tags(*texts):
    t = " ".join(x for x in texts if x)
    tags = []
    if re.search(r"vis|VNIR", t, re.I):
        tags.append("vis-NIR")
    if re.search(r"\bMIR\b|mid-?infrared|mid-IR|DRIFT|FTIR", t, re.I):
        tags.append("MIR")
    if re.search(r"LWIR|thermal|\bTIR\b", t, re.I):
        tags.append("LWIR / thermal")
    if not tags and re.search(r"NIR", t):
        tags.append("NIR")
    return tags or ["Not recorded"]


def short_ref(authors, year):
    if not authors:
        return f"Unknown {year or ''}".strip()
    a = str(authors)
    if a.startswith("Agriculture and Agri-Food"):
        return "AAFC CPSSL web resource"
    parts = [p.strip() for p in a.split(";")]
    etal = any(p.lower().startswith("et al") for p in parts)
    parts = [p for p in parts if p and not p.lower().startswith("et al")]

    def surname(p):
        out = []
        for tok in p.split(" "):
            if re.match(r"^([A-Z]\.-?)+[A-Z]?\.?$", tok):
                break
            out.append(tok)
        return " ".join(out) or p

    if len(parts) == 1 and not etal:
        return f"{surname(parts[0])} {year}"
    if len(parts) == 2 and not etal:
        return f"{surname(parts[0])} & {surname(parts[1])} {year}"
    return f"{surname(parts[0])} et al. {year}"


def ring_area_centroid(ring):
    a = cx = cy = 0.0
    for (x0, y0), (x1, y1) in zip(ring, ring[1:] + ring[:1]):
        f = x0 * y1 - x1 * y0
        a += f; cx += (x0 + x1) * f; cy += (y0 + y1) * f
    if abs(a) < 1e-12:
        return 0.0, ring[0]
    return a / 2, (cx / (3 * a), cy / (3 * a))


def label_point(geom):
    """Centroid of the largest polygon part (keeps e.g. France on the mainland)."""
    polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
    best = max((ring_area_centroid(p[0]) for p in polys), key=lambda t: abs(t[0]))
    return [round(best[1][0], 3), round(best[1][1], 3)]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    for name in ("Papers", "Libraries", "Library countries"):
        if name not in wb.sheetnames:
            sys.exit(f"Sheet '{name}' not found in {SRC}")

    # ---- papers ----
    ws = wb["Papers"]; h = header_map(ws)
    def col(row, name):
        return row[h[name]] if name in h else None
    papers = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not col(row, "Title"):
            continue
        c1, c2, c3 = (str(col(row, k) or "").strip() for k in ("C1 Soil-specific", "C2 Spectroscopic", "C3 Library-scale data"))
        dup = col(row, "Duplicate of row")
        if dup not in (None, ""):
            decision = "Exclude (duplicate)"
        elif "Check" in (c1, c2, c3):
            decision = "Check"
        elif (c1, c2, c3) == ("Yes", "Yes", "Yes"):
            decision = "Include"
        else:
            decision = "Exclude"
        doi = clean(col(row, "DOI"))
        if doi and not re.match(r"^10\.\S+/\S+", doi):
            doi = None
        papers.append({
            "row": i,
            "ref": short_ref(col(row, "Authors"), col(row, "Year")),
            "authors": clean(col(row, "Authors")),
            "title": clean(col(row, "Title")),
            "year": clean(col(row, "Year")),
            "doi": doi,
            "decision": decision,
            "role": clean(col(row, "Paper role")) if col(row, "Paper role") != "—" else None,
            "libraries": [x for x in (clean(col(row, "Library ID")), clean(col(row, "Secondary library ID"))) if x],
            "note": clean(col(row, "Exclusion reason / notes")),
            "samples": clean(col(row, "Number of Samples")),
            "spectral": clean(col(row, "Spectral Region")),
            "country_text": clean(col(row, "Country/Region")),
        })

    # ---- library countries ----
    lc = wb["Library countries"]; hc = header_map(lc)
    lib_countries = {}
    for row in lc.iter_rows(min_row=2, values_only=True):
        lid, cname = row[hc["Library ID"]], row[hc["Country"]]
        if lid and cname:
            lib_countries.setdefault(lid, []).append(cname)

    # ---- world boundaries ----
    world = json.loads(WORLD.read_text())
    by_name = {f["properties"]["name"]: f for f in world["features"]}
    missing = {c for cs in lib_countries.values() for c in cs if ALIAS.get(c, c) not in by_name}
    if missing:
        sys.exit(f"Countries not found in boundary file (add to ALIAS): {sorted(missing)}")
    centroid = {c: label_point(by_name[ALIAS.get(c, c)]["geometry"]) for cs in lib_countries.values() for c in cs}

    # ---- libraries ----
    lb = wb["Libraries"]; hl = header_map(lb)
    size_col = next(k for k in hl if k.startswith("Largest size") or k.startswith("Reported size"))
    libs = []
    for row in lb.iter_rows(min_row=2, values_only=True):
        lid = row[hl["Library ID"]]
        if not lid:
            continue
        linked = [p for p in papers if lid in p["libraries"] and p["decision"] == "Include"]
        intro = next((p for p in linked if p["role"] == "Introduces" and p["libraries"][0] == lid), None)
        countries = lib_countries.get(lid, [])
        name = row[hl["Library name"]]
        spec = clean(row[hl["Spectral region"]])
        samples = row[hl["Samples (numeric)"]]
        try:
            samples = int(samples) if samples not in (None, "") else None
        except (TypeError, ValueError):
            samples = None
        pt = None
        lon = row[hl["Map longitude (approx.)"]] if "Map longitude (approx.)" in hl else None
        lat = row[hl["Map latitude (approx.)"]] if "Map latitude (approx.)" in hl else None
        if isinstance(lon, (int, float)) and isinstance(lat, (int, float)):
            pt = [round(float(lon), 3), round(float(lat), 3)]
        elif countries:
            xs = [centroid[c][0] for c in countries]; ys = [centroid[c][1] for c in countries]
            pt = [round(sum(xs) / len(xs), 3), round(sum(ys) / len(ys), 3)]
        libs.append({
            "id": lid,
            "name": name,
            "scale": row[hl["Scale"]],
            "countries": countries,
            "continent": row[hl["Continent"]],
            "spectral": spec,
            "spectral_tags": spectral_tags(spec, name, clean(row[hl[size_col]]), *[p["spectral"] for p in linked]),
            "samples": samples,
            "size_text": clean(row[hl[size_col]]),
            "availability": clean(row[hl["Data availability"]]),
            "intro_in_review": intro is not None,
            "intro_ref": intro["ref"] if intro else None,
            "intro_row": intro["row"] if intro else None,
            "papers": [p["row"] for p in linked],
            "notes": clean(row[hl["Notes"]]),
            "point": pt,  # [lon, lat]
            "point_source": ("study region" if isinstance(lon, (int, float)) else "country centroid") if pt else None,
        })
    lib_ids = {l["id"] for l in libs}
    for p in papers:
        bad = [x for x in p["libraries"] if x not in lib_ids]
        if bad:
            print(f"warning: Papers row {p['row']} links unknown library {bad}")

    # ---- country GeoJSON with coverage attributes ----
    cov = {}
    for l in libs:
        if not l["papers"]:
            continue
        for c in l["countries"]:
            cov.setdefault(ALIAS.get(c, c), []).append(l)
    feats = []
    for f in world["features"]:
        ls = cov.get(f["properties"]["name"], [])
        best = min((SCALE_RANK[l["scale"]] for l in ls if l["scale"] in SCALE_RANK), default=None)
        props = {
            "name": f["properties"]["name"],
            "iso_n3": f["properties"].get("iso_n3"),
            "n_libraries": len(ls),
            "n_national": sum(l["scale"] == "National" for l in ls),
            "n_subnational": sum(l["scale"] == "Sub-national" for l in ls),
            "n_multicountry": sum(l["scale"] == "Multi-country" for l in ls),
            "highest_scale": {0: "National", 1: "Sub-national only", 2: "Multi-country only"}.get(best, "None"),
            "library_ids": ";".join(l["id"] for l in ls),
        }
        feats.append({"type": "Feature", "properties": props, "geometry": f["geometry"]})

    OUT.joinpath("papers.json").write_text(json.dumps(papers, ensure_ascii=False, indent=0))
    OUT.joinpath("libraries.json").write_text(json.dumps(libs, ensure_ascii=False, indent=0))
    OUT.joinpath("countries.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": feats}, ensure_ascii=False, separators=(",", ":")))
    meta = {
        "built": datetime.date.today().isoformat(),
        "source": SRC.name,
        "papers": len(papers),
        "papers_included": sum(p["decision"] == "Include" for p in papers),
        "libraries": len(libs),
        "countries": sum(1 for f in feats if f["properties"]["n_libraries"] > 0),
    }
    OUT.joinpath("meta.json").write_text(json.dumps(meta, indent=1))
    print(json.dumps(meta))


if __name__ == "__main__":
    main()
